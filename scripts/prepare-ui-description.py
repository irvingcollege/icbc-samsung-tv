#!/usr/bin/env python3
"""Fill Samsung v1.4.2 UI template with ICBC content and genuine screenshots."""
import argparse
from copy import deepcopy
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

ROOT=Path(__file__).resolve().parent.parent
parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('template',type=Path)
parser.add_argument('--output',type=Path,default=ROOT/'downloads/ICBC-TV-Application-UI-Description-v1.0.1.pptx')
args=parser.parse_args()
p=Presentation(args.template);original=list(p.slides)
if len(original)!=21:raise RuntimeError('Expected Samsung v1.4.2 21-slide template; inspect before modifying another layout')
shots=ROOT/'artifacts/seller-office-assets'

def text(shape,value,size=18):
    shape.text=value
    for para in shape.text_frame.paragraphs:
        para.space_after=Pt(9)
        for run in para.runs:
            run.font.name='Arial';run.font.size=Pt(size);run.font.color.rgb=RGBColor(25,32,36)
    shape.text_frame.word_wrap=True

def box(slide,value,x,y,w,h,size=18):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));text(shape,value,size);return shape

def clear(slide):
    for shape in list(slide.shapes)[2:]:shape._element.getparent().remove(shape._element)

def title(slide,value):text(slide.shapes[0],value,26)

def body(slide,value,size=19):
    clear(slide);box(slide,value,.55,1.55,8.9,5.45,size)

def table(slide,rows,x=.55,y=2,w=8.9,h=4.5,widths=None,size=12):
    t=slide.shapes.add_table(len(rows),len(rows[0]), Inches(x),Inches(y),Inches(w),Inches(h)).table
    if widths:
        for c,width in zip(t.columns,widths):c.width=Inches(width)
    for i,row in enumerate(rows):
        for j,value in enumerate(row):
            c=t.cell(i,j);c.text=value
            for para in c.text_frame.paragraphs:
                for run in para.runs:run.font.name='Arial';run.font.size=Pt(size);run.font.color.rgb=RGBColor(20,28,32)
    return t

def preview(slide,subtitle,image,rows,caption):
    clear(slide);title(slide,'Menus and Functions')
    box(slide,subtitle,.55,1.13,8.9,.45,18)
    slide.shapes.add_picture(str(shots/image),Inches(.55),Inches(1.8),width=Inches(6.1))
    table(slide,[['UI element','Function']]+rows,x=6.82,y=1.8,w=2.65,h=3.95,widths=[1,1.65],size=10)
    box(slide,caption,.55,6.0,8.9,.9,12)

# Cover: fill the template's existing empty title fields.
text(original[1].shapes[0],'Irving College Baptist Church',30)
text(original[1].shapes[1],'Application UI Description | App version 1.0.1\nSeller Office App ID: 3202610049560',16)
# Revision history retains the template's table and columns.
rev=original[2].shapes[2].table
values=['1.1','2026-10-18','Added actual app preview images and completed navigation, use cases, menus and key mappings. Documentation correction for resubmission.','Irving College Baptist Church','1.0.1']
for j,value in enumerate(values):rev.cell(1,j).text=value
rev.rows[1].height=Inches(1.3)
for row in list(rev.rows)[2:]:
    for cell in row.cells:cell.text=''
for row in rev.rows:
    for cell in row.cells:
        for para in cell.text_frame.paragraphs:
            for run in para.runs:run.font.name='Arial';run.font.size=Pt(12);run.font.color.rgb=RGBColor(20,28,32)
for index in [3,8,13,17,19]:
    text(original[index].shapes[2],'UI Structure\nUse Cases\nMenus and Functions — App preview images\nRemote Control Keys\nLanguage Options',23)

s=original[4];clear(s);title(s,'UI Structure')
nodes=[('Home: church info + service times',2.5,1.5,5,0.6),('Watch Live / HLS',.65,2.8,3.7,.6),('Browse Sermons / Refresh',5.2,2.8,4.1,.6),('Player: live or archived MP4',2.4,4.2,5.2,.7),('Home → RETURN → Exit confirmation',2.0,5.8,6,.6)]
for label,x,y,w,h in nodes:
    shape=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));shape.fill.solid();shape.fill.fore_color.rgb=RGBColor(233,241,240);shape.line.color.rgb=RGBColor(36,75,71);text(shape,label,15)
for x1,y1,x2,y2 in [(4,2.1,2.5,2.8),(6,2.1,7.2,2.8),(2.5,3.4,4,4.2),(7.2,3.4,6,4.2)]:
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));c.line.color.rgb=RGBColor(36,75,71)
box(s,'Select an action or sermon with OK. Player Back/RETURN restores Home and prior focus.\nErrors remain on the relevant screen: Retry for playback; Refresh for the catalog.\nExit confirmation: Stay returns Home; Exit App closes the application.',.6,6.45,8.8,.65,11)

body(original[9],'Main use cases\n• Watch the public live HLS stream when a broadcast is available.\n• Browse the public sermon catalog and play an archived MP4 sermon.\n• Read church branding, website and Central Time service schedule on Home.\n• Use TV remote arrows/OK, Return and media controls.\n• Recover from unavailable streams, catalog failures and empty catalogs.\n• Confirm application exit from Home.\nAccess: free; no login, subscription, purchasing or test credentials.',19)
title(original[10],'Use Cases — Watch Live')
body(original[10],'1. On Home, focus Watch Live and press OK.\n2. The full-screen player connects to the public HLS broadcast.\n3. On a Samsung TV the app uses native AVPlay; browser previews use HTML video/HLS.js.\n4. Use Play/Pause or the on-screen button. Live rewind/forward controls are disabled.\n5. If the stream is offline or fails, read the status message and choose Retry.\n6. Choose Back to Home or press RETURN to close playback and restore prior focus.\n\nStream: https://live.irvingcollege.church/hls/stream.m3u8\nBroadcast availability is not guaranteed outside church services. No successful physical-TV playback is claimed in this document.',17)
title(original[11],'Use Cases — Browse Sermons')
body(original[11],'1. On Home select Browse Sermons; focus moves to the first available sermon (or Refresh for an empty catalog).\n2. Use arrows to select a sermon card; press OK to open its archived MP4.\n3. In the player use Play/Pause, −30 seconds and +30 seconds. Dedicated media keys provide the same functions.\n4. RETURN / Back to Home closes the player and restores the selected sermon.\n5. Refresh requests the public catalog again. An empty list has a clear message; network/format errors explain recovery.\n\nCatalog: https://live.irvingcollege.church/sermons/catalog.json\nNo test account is required. Catalog titles and dates are provided by the church server.',17)
title(original[12],'Use Cases — Access, Errors and Exit')
body(original[12],'No login or in-app purchase\nAll content is publicly accessible. No account/password, payment, subscription, advertisement or purchase flow is implemented.\n\nNetwork recovery\nLoading and unavailable-stream messages appear in the player. Retry reconnects the selected video. The catalog provides Refresh and distinct empty/error messages. No fake replacement videos are shown.\n\nExit and background behavior\nRETURN closes playback, returns from archive scrolling to Home, or opens the Home exit confirmation. Stay dismisses the confirmation; Exit App exits. Hidden playback is paused. The platform EXIT key remains under Samsung system behavior.',17)
preview(original[14],'1.- App preview images — Home','01-home.jpg',[['Watch Live','Open live player'],['Browse Sermons','Focus archive'],['Service times','Central Time'],['Refresh','Reload catalog'],['Sermon card','Open MP4']], 'Genuine Chromium rendering of the app. Existing app artwork is shown unchanged. Catalog content is the real church response captured on 2026-10-13; this is not a physical-TV photograph.')
preview(original[15],'App preview images — Sermon browsing','02-sermons.jpg',[['White outline','Current focus'],['Title / date','Catalog metadata'],['Sermon card','OK starts video'],['Arrows','Move focus'],['Refresh','Retry catalog']], 'Genuine app screenshot: archive focus state. No fabricated sermon records or video imagery. Browser capture used unmodified app files and unchanged real catalog bytes through the capture harness.')
preview(original[16],'App preview images — Exit confirmation','03-exit-confirmation.jpg',[['Stay','Return to Home'],['Exit App','Close app'],['RETURN','Dismiss dialog'],['Arrows / OK','Choose / activate']], 'Genuine browser-rendered exit dialog from the app. No login or payment is required. This screen is reached with RETURN from Home.')
# One extra screenshot slide is allowed by the template's Menus and Functions instructions.
player=p.slides.add_slide(original[14].slide_layout)
for shape in list(player.shapes):shape._element.getparent().remove(shape._element)
for shape in list(original[14].shapes)[:2]:player.shapes._spTree.insert_element_before(deepcopy(shape._element),'p:extLst')
preview(player,'App preview images — Player and network recovery','04-player-browser.jpg',[['Pause / Play','Toggle playback'],['−30 / +30 sec','Archive seek only'],['Retry','Reconnect video'],['Back to Home','Stop / return'],['Status text','Connection/error']], 'Genuine browser capture of the player’s playback-error state on 2026-10-18. No successful playback or video frame is fabricated. Samsung AVPlay, live broadcast availability and physical-TV playback remain unverified.')

s=original[18];mandatory=[[c.text for c in row.cells] for row in s.shapes[3].table.rows][:3];clear(s);title(s,'Remote Control Keys')
box(s,'RETURN: closes player/restores focus; on Home, opens exit confirmation.\nEXIT: platform-defined behavior; the app does not register or override this key.',.55,1.15,8.9,.75,13)
rows=mandatory+[['ENTER / OK','Activate the focused button or sermon','No full-screen toggle mapping'],['UP/DOWN/LEFT/RIGHT','Move focus between visible controls','Scroll selected archive item into view'],['PLAY / PAUSE / PLAY-PAUSE','Control active video','Registered optional media keys'],['STOP','Close player and return Home','While playing'],['REWIND / FAST FORWARD','Seek −30 / +30 seconds','Archived content only'],['CHANNEL / VOLUME / COLORS','No custom app mapping','System behavior retained']]
table(s,rows,y=2.05,h=4.65,widths=[1.7,4.5,2.7],size=11)
title(original[20],'Language Options')
body(original[20],'The application does not support multiple languages.\n\nThe application interface is English. There is no in-app language selector. Sermon titles, speaker fields and date strings are displayed as supplied by the public catalog.\n\nChurch service times are Central Time:\nSunday School 9:30 AM; Sunday Worship 10:30 AM.\nWednesday Meal 5:30 PM; Bible Study 6:30 PM.',20)

# Delete the instruction slide and unused example diagram slides as explicitly instructed.
order=[slide for i,slide in enumerate(original) if i not in [0,5,6,7]]
order.insert(order.index(original[17]),player)
ids={id(slide):next(sid for sid in p.slides._sldIdLst if p.part.related_slide(sid.rId)==slide) for slide in list(p.slides)}
for slide in list(p.slides):
    if slide not in order:p.part.drop_rel(ids[id(slide)].rId)
for sid in list(p.slides._sldIdLst):p.slides._sldIdLst.remove(sid)
for i,slide in enumerate(order):
    p.slides._sldIdLst.append(ids[id(slide)])
    if len(slide.shapes)>1 and slide.shapes[1].has_text_frame and slide.shapes[1].width<Inches(3):text(slide.shapes[1],str(i+1),10)
p.core_properties.title='ICBC Application UI Description — 1.0.1'
p.core_properties.author='Irving College Baptist Church'
p.core_properties.subject='Samsung TV review resubmission: actual app preview images'
args.output.parent.mkdir(parents=True,exist_ok=True);p.save(args.output)
print(args.output,'slides:',len(p.slides))
